"""
RAG pipeline: document upload → chunk → embed → ChromaDB Cloud.
Query retrieves relevant chunks to augment LLM context.
LLM: Sarvam AI Sarvam-M (free chat completion). See https://docs.sarvam.ai
Env: CHROMA_*, SARVAM_API_KEY (from https://dashboard.sarvam.ai).
"""
import logging
import os
import re
import sys
import uuid
from pathlib import Path

from dotenv import load_dotenv
load_dotenv()

# ChromaDB imported lazily via _import_chromadb() — not compatible with Python 3.14 (Pydantic v1).
from pypdf import PdfReader
from docx import Document as DocxDocument
from docx import Document as DocxDocument

log = logging.getLogger(__name__)

# ChromaDB Cloud credentials from .env
CHROMA_API_KEY = (os.getenv("CHROMA_API_KEY") or "").strip()
CHROMA_TENANT = (os.getenv("CHROMA_TENANT") or "").strip()
CHROMA_DATABASE = (os.getenv("CHROMA_DATABASE") or "").strip()
COLLECTION_NAME = "caller_rag"

# Sarvam AI (Sarvam-M chat completion is free; get key from https://dashboard.sarvam.ai)
SARVAM_API_KEY = (os.getenv("SARVAM_API_KEY") or os.getenv("api_subscription_key") or "").strip()
SARVAM_MODEL = "sarvam-m"

# Chunking defaults
CHUNK_SIZE = 512
CHUNK_OVERLAP = 80
QUERY_N_RESULTS = 10


def _import_chromadb():
    """Import chromadb; on Python 3.14 this can raise due to Pydantic v1 incompatibility."""
    if sys.version_info >= (3, 14):
        try:
            import chromadb
            return chromadb
        except Exception as e:
            raise RuntimeError(
                "ChromaDB is not compatible with Python 3.14 (Pydantic v1). "
                "Use Python 3.12 or 3.13 for RAG, e.g.: pyenv install 3.13 && pyenv local 3.13"
            ) from e
    import chromadb
    return chromadb


def _get_embedding_function():
    """Sentence-transformers embedding for Chroma (runs locally)."""
    chromadb = _import_chromadb()
    from chromadb.utils import embedding_functions
    return embedding_functions.SentenceTransformerEmbeddingFunction(
        model_name="all-MiniLM-L6-v2"
    )


def _get_client():
    if not CHROMA_API_KEY or not CHROMA_TENANT or not CHROMA_DATABASE:
        raise ValueError(
            "CHROMA_API_KEY, CHROMA_TENANT, and CHROMA_DATABASE must be set in .env"
        )
    chromadb = _import_chromadb()
    return chromadb.CloudClient(
        tenant=CHROMA_TENANT,
        database=CHROMA_DATABASE,
        api_key=CHROMA_API_KEY,
    )


def _get_collection():
    client = _get_client()
    return client.get_or_create_collection(
        name=COLLECTION_NAME,
        embedding_function=_get_embedding_function(),
    )


def extract_text_from_file(file_path: str) -> str:
    """Extract plain text from PDF or DOCX. Raises ValueError for unsupported type."""
    path = Path(file_path)
    suffix = path.suffix.lower()
    text = ""

    if suffix == ".pdf":
        reader = PdfReader(file_path)
        for page in reader.pages:
            text += (page.extract_text() or "") + "\n"
    elif suffix in (".docx", ".doc"):
        doc = DocxDocument(file_path)
        for para in doc.paragraphs:
            text += para.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + " "
                text += "\n"
    else:
        raise ValueError(f"Unsupported file type: {suffix}. Use .pdf or .docx")

    return re.sub(r"\n{3,}", "\n\n", text).strip()


def chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Split text into overlapping chunks for embedding."""
    if not text or not text.strip():
        return []
    chunks = []
    start = 0
    text = text.replace("\r\n", "\n")
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end]
        if end < len(text):
            last_space = chunk.rfind(" ")
            if last_space > chunk_size // 2:
                end = start + last_space + 1
                chunk = text[start:end]
        chunk = chunk.strip()
        if chunk:
            chunks.append(chunk)
        start = end - overlap
        if start < 0:
            start = end
    return chunks


def ingest_document(file_path: str) -> int:
    """
    Load a PDF or DOCX, chunk it, and add to ChromaDB.
    Returns number of chunks added.
    """
    text = extract_text_from_file(file_path)
    chunks = chunk_text(text)
    if not chunks:
        log.warning("No text chunks from %s", file_path)
        return 0

    coll = _get_collection()
    ids = [str(uuid.uuid4()) for _ in chunks]
    coll.add(ids=ids, documents=chunks)
    log.info("RAG: ingested %s → %d chunks", file_path, len(chunks))
    return len(chunks)


def query(query_text: str, n_results: int = QUERY_N_RESULTS) -> str:
    """
    Retrieve relevant chunks from ChromaDB for the query.
    Returns a single string of concatenated chunks, or empty if none/collection empty.
    """ 

    if not query_text or not query_text.strip():
        return ""
    try:
        coll = _get_collection()
        if coll.count() == 0:
            return ""
        result = coll.query(
            query_texts=[query_text.strip()],
            n_results=min(n_results, coll.count()),
            include=["documents"],
        )
        docs = result.get("documents") or []
        if not docs or not docs[0]:
            return ""
        return "\n\n".join(docs[0]).strip()
    except Exception as e:
        log.warning("RAG query failed: %s", e)
        return ""


def query_opening_and_general(user_query: str, opening_n: int = 4, general_n: int = 6) -> str:
    """
    Retrieve context for voice agent: opening/greeting from document + content relevant to user.
    Ensures the agent gets the document's opening script (e.g. नमस्ते, मैं सोनी...) and user-relevant info.
    """
    opening = query(
        "opening greeting introduction first message namaste hello company name assistant",
        n_results=opening_n,
    )
    general = query(user_query or "company services policies enquiry", n_results=general_n)
    parts = [p for p in [opening, general] if p]
    if not parts:
        return ""
    return "\n\n".join(parts).strip()


def get_document_count() -> int:
    """Return number of chunks (documents) in the RAG collection."""
    try:
        return _get_collection().count()
    except Exception:
        return 0


# --- Sarvam AI LLM (Sarvam-M; free chat completion per official docs) ---

def _get_sarvam_client():
    """Return SarvamAI client if SARVAM_API_KEY is set, else None."""
    if not SARVAM_API_KEY:
        return None
    try:
        from sarvamai import SarvamAI
        return SarvamAI(api_subscription_key=SARVAM_API_KEY)
    except ImportError:
        log.warning("sarvamai not installed; pip install sarvamai")
        return None


def sarvam_chat(
    messages: list[dict],
    *,
    model: str = SARVAM_MODEL,
    temperature: float = 0.2,
    max_tokens: int | None = 2048,
    reasoning_effort: str | None = "low",
) -> str:
    """
    Call Sarvam AI chat completion (Sarvam-M, free).
    messages: list of {"role": "user"|"assistant"|"system", "content": "..."}
    Returns the assistant reply text, or empty string on error.

    For low latency (e.g. voice): use max_tokens=100–150 and reasoning_effort="low"
    (Sarvam docs: fewer tokens = faster; low effort = quick replies, less thinking).
    """
    client = _get_sarvam_client()
    if not client:
        log.warning("SARVAM_API_KEY not set; cannot call Sarvam LLM")
        return ""
    try:
        kwargs = {
            "messages": messages,
            "model": model,
            "temperature": temperature,
        }
        if max_tokens is not None:
            kwargs["max_tokens"] = max_tokens
        if reasoning_effort is not None:
            kwargs["reasoning_effort"] = reasoning_effort
        response = client.chat.completions(**kwargs)
        choices = getattr(response, "choices", None) or []
        if not choices:
            return ""
        msg = choices[0].message if hasattr(choices[0], "message") else choices[0].get("message")
        content = getattr(msg, "content", None) or (msg.get("content") if isinstance(msg, dict) else None)
        return (content or "").strip()
    except Exception as e:
        log.warning("Sarvam chat failed: %s", e)
        return ""


def answer_with_rag(user_query: str) -> str:
    """
    Get RAG context (opening + user-relevant chunks) then answer using Sarvam-M.
    Uses query_opening_and_general for context. Returns LLM reply or empty on error.
    """
    context = query_opening_and_general(user_query)
    if context:
        messages = [
            {"role": "system", "content": "Use the following document context to answer the user. Be concise and accurate.\n\n" + context},
            {"role": "user", "content": user_query or "What can you help with?"},
        ]
    else:
        messages = [{"role": "user", "content": user_query or "Hello."}]
    return sarvam_chat(messages)
